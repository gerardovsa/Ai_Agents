"""
Test Markdown Text Extraction
==============================

Test the enhanced text extractor with Markdown formatting
for DOCX, XLSX, PPTX, CSV, JSON files.
"""

import sys
import os

# Add AI_infrastructure to path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'AI_infrastructure'))

from AI_infrastructure.core.text_extractor import TextExtractor


def test_docx_markdown():
    """Test DOCX extraction with Markdown formatting"""
    print("\n" + "="*70)
    print("TEST 1: DOCX with Markdown Formatting")
    print("="*70)
    
    # Create a sample DOCX in memory
    try:
        from docx import Document
        from io import BytesIO
        
        doc = Document()
        doc.add_heading('Business Proposal', 0)
        doc.add_heading('Executive Summary', level=1)
        doc.add_paragraph('This is a bold and important statement.', style='Normal')
        
        # Add a paragraph with formatting
        para = doc.add_paragraph()
        para.add_run('Normal text, ').bold = False
        para.add_run('bold text, ').bold = True
        para.add_run('italic text, ').italic = True
        para.add_run('bold+italic text').bold = True
        para.runs[-1].italic = True
        
        doc.add_heading('Key Features', level=2)
        doc.add_paragraph('Feature 1: Fast processing', style='List Bullet')
        doc.add_paragraph('Feature 2: Cost effective', style='List Bullet')
        doc.add_paragraph('Feature 3: Scalable', style='List Bullet')
        
        # Add a table
        table = doc.add_table(rows=3, cols=3)
        table.style = 'Light Grid Accent 1'
        
        # Headers
        table.rows[0].cells[0].text = 'Product'
        table.rows[0].cells[1].text = 'Price'
        table.rows[0].cells[2].text = 'Quantity'
        
        # Data
        table.rows[1].cells[0].text = 'Widget A'
        table.rows[1].cells[1].text = '$50'
        table.rows[1].cells[2].text = '100'
        
        table.rows[2].cells[0].text = 'Widget B'
        table.rows[2].cells[1].text = '$75'
        table.rows[2].cells[2].text = '50'
        
        # Save to BytesIO
        docx_bytes = BytesIO()
        doc.save(docx_bytes)
        docx_bytes.seek(0)
        
        # Extract text
        extractor = TextExtractor()
        result = extractor.extract_text(
            file_data=docx_bytes.read(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            filename='business_proposal.docx',
            output_format='markdown'
        )
        
        if result['success']:
            print("\n[OK] DOCX Extraction Successful!")
            print(f"\nFormat: {result.get('format', 'N/A')}")
            print(f"Word Count: {result['metadata']['word_count']}")
            print("\n--- Extracted Markdown Content ---\n")
            print(result['text'])
        else:
            print(f"\n[FAIL] DOCX Extraction Failed: {result.get('error')}")
            
    except ImportError as e:
        print(f"\n⚠️ Skipping DOCX test - Missing library: {e}")


def test_xlsx_markdown():
    """Test XLSX extraction with Markdown table formatting"""
    print("\n" + "="*70)
    print("TEST 2: XLSX with Markdown Table Formatting")
    print("="*70)
    
    try:
        from openpyxl import Workbook
        from io import BytesIO
        
        wb = Workbook()
        
        # Sheet 1: Sales Data
        ws1 = wb.active
        ws1.title = "Sales Q4 2025"
        ws1.append(['Month', 'Revenue', 'Profit', 'Growth'])
        ws1.append(['October', '$150,000', '$45,000', '12%'])
        ws1.append(['November', '$175,000', '$52,500', '16%'])
        ws1.append(['December', '$200,000', '$60,000', '14%'])
        
        # Sheet 2: Employee Data
        ws2 = wb.create_sheet("Employees")
        ws2.append(['Name', 'Department', 'Salary', 'Start Date'])
        ws2.append(['John Smith', 'Engineering', '$120,000', '2023-01-15'])
        ws2.append(['Jane Doe', 'Marketing', '$95,000', '2023-03-20'])
        ws2.append(['Bob Johnson', 'Sales', '$110,000', '2022-11-05'])
        
        # Save to BytesIO
        xlsx_bytes = BytesIO()
        wb.save(xlsx_bytes)
        xlsx_bytes.seek(0)
        
        # Extract text
        extractor = TextExtractor()
        result = extractor.extract_text(
            file_data=xlsx_bytes.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            filename='company_data.xlsx',
            output_format='markdown'
        )
        
        if result['success']:
            print("\n✅ XLSX Extraction Successful!")
            print(f"\nFormat: {result.get('format', 'N/A')}")
            print(f"Word Count: {result['metadata']['word_count']}")
            print("\n--- Extracted Markdown Content ---\n")
            print(result['text'])
        else:
            print(f"\n❌ XLSX Extraction Failed: {result.get('error')}")
            
    except ImportError as e:
        print(f"\n⚠️ Skipping XLSX test - Missing library: {e}")


def test_pptx_markdown():
    """Test PPTX extraction with Markdown formatting"""
    print("\n" + "="*70)
    print("TEST 3: PPTX with Markdown Formatting")
    print("="*70)
    
    try:
        from pptx import Presentation
        from pptx.util import Inches
        from io import BytesIO
        
        prs = Presentation()
        
        # Slide 1: Title Slide
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "Q4 2025 Business Review"
        slide1.placeholders[1].text = "Prepared by Finance Team"
        
        # Slide 2: Bullet Points
        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        slide2.shapes.title.text = "Key Achievements"
        
        content = slide2.placeholders[1].text_frame
        content.text = "Revenue Growth\nCustomer Acquisition\nProduct Launch\nMarket Expansion"
        
        # Slide 3: Summary
        slide3 = prs.slides.add_slide(prs.slide_layouts[1])
        slide3.shapes.title.text = "Next Quarter Goals"
        
        goals = slide3.placeholders[1].text_frame
        goals.text = "Increase sales by 20%\nLaunch new product line\nExpand to 3 new markets\nImprove customer satisfaction"
        
        # Save to BytesIO
        pptx_bytes = BytesIO()
        prs.save(pptx_bytes)
        pptx_bytes.seek(0)
        
        # Extract text
        extractor = TextExtractor()
        result = extractor.extract_text(
            file_data=pptx_bytes.read(),
            content_type='application/vnd.openxmlformats-officedocument.presentationml.presentation',
            filename='business_review.pptx',
            output_format='markdown'
        )
        
        if result['success']:
            print("\n✅ PPTX Extraction Successful!")
            print(f"\nFormat: {result.get('format', 'N/A')}")
            print(f"Word Count: {result['metadata']['word_count']}")
            print("\n--- Extracted Markdown Content ---\n")
            print(result['text'])
        else:
            print(f"\n❌ PPTX Extraction Failed: {result.get('error')}")
            
    except ImportError as e:
        print(f"\n⚠️ Skipping PPTX test - Missing library: {e}")


def test_csv_markdown():
    """Test CSV extraction with Markdown table formatting"""
    print("\n" + "="*70)
    print("TEST 4: CSV with Markdown Table Formatting")
    print("="*70)
    
    csv_content = """Product,Price,Stock,Category
Laptop,$1200,45,Electronics
Mouse,$25,150,Accessories
Keyboard,$75,80,Accessories
Monitor,$350,60,Electronics
Headphones,$100,120,Accessories"""
    
    extractor = TextExtractor()
    result = extractor.extract_text(
        file_data=csv_content.encode('utf-8'),
        content_type='text/csv',
        filename='inventory.csv',
        output_format='markdown'
    )
    
    if result['success']:
        print("\n✅ CSV Extraction Successful!")
        print(f"\nFormat: {result.get('format', 'N/A')}")
        print(f"Word Count: {result['metadata']['word_count']}")
        print("\n--- Extracted Markdown Content ---\n")
        print(result['text'])
    else:
        print(f"\n❌ CSV Extraction Failed: {result.get('error')}")


def test_json_markdown():
    """Test JSON extraction with Markdown code block"""
    print("\n" + "="*70)
    print("TEST 5: JSON with Markdown Code Block")
    print("="*70)
    
    json_content = """{
  "company": "InHouse Print Solutions",
  "employees": 50,
  "departments": [
    {"name": "Sales", "headcount": 12},
    {"name": "Production", "headcount": 25},
    {"name": "Admin", "headcount": 8},
    {"name": "IT", "headcount": 5}
  ],
  "annual_revenue": 2500000,
  "established": 2010
}"""
    
    extractor = TextExtractor()
    result = extractor.extract_text(
        file_data=json_content.encode('utf-8'),
        content_type='application/json',
        filename='company_data.json',
        output_format='markdown'
    )
    
    if result['success']:
        print("\n✅ JSON Extraction Successful!")
        print(f"\nFormat: {result.get('format', 'N/A')}")
        print(f"Word Count: {result['metadata']['word_count']}")
        print("\n--- Extracted Markdown Content ---\n")
        print(result['text'])
    else:
        print(f"\n❌ JSON Extraction Failed: {result.get('error')}")


def test_python_code_markdown():
    """Test Python code extraction with Markdown code block"""
    print("\n" + "="*70)
    print("TEST 6: Python Code with Markdown Code Block")
    print("="*70)
    
    python_content = """def calculate_total(items):
    \"\"\"Calculate total price of items\"\"\"
    total = 0
    for item in items:
        total += item['price'] * item['quantity']
    return total

# Test data
items = [
    {'name': 'Widget', 'price': 50, 'quantity': 10},
    {'name': 'Gadget', 'price': 75, 'quantity': 5}
]

result = calculate_total(items)
print(f'Total: ${result}')
"""
    
    extractor = TextExtractor()
    result = extractor.extract_text(
        file_data=python_content.encode('utf-8'),
        content_type='text/x-python',
        filename='calculator.py',
        output_format='markdown'
    )
    
    if result['success']:
        print("\n✅ Python Code Extraction Successful!")
        print(f"\nFormat: {result.get('format', 'N/A')}")
        print(f"Word Count: {result['metadata']['word_count']}")
        print("\n--- Extracted Markdown Content ---\n")
        print(result['text'])
    else:
        print(f"\n❌ Python Code Extraction Failed: {result.get('error')}")


def main():
    """Run all tests"""
    print("\n" + "="*70)
    print("MARKDOWN TEXT EXTRACTION TEST SUITE")
    print("Testing enhanced text extractor with Markdown formatting")
    print("="*70)
    
    # Run tests
    test_docx_markdown()
    test_xlsx_markdown()
    test_pptx_markdown()
    test_csv_markdown()
    test_json_markdown()
    test_python_code_markdown()
    
    print("\n" + "="*70)
    print("ALL TESTS COMPLETED!")
    print("="*70)
    print("\n✅ Features Demonstrated:")
    print("  - DOCX → Markdown (headings, bold, italic, lists, tables)")
    print("  - XLSX → Markdown tables (multi-sheet support)")
    print("  - PPTX → Structured slides with bullet points")
    print("  - CSV → Markdown tables")
    print("  - JSON → Code blocks with syntax highlighting")
    print("  - Python → Code blocks with language tags")
    print("\n🎯 Benefits:")
    print("  - Preserves document structure")
    print("  - Better AI comprehension")
    print("  - Readable formatting")
    print("  - No raw base64 in chat")


if __name__ == '__main__':
    main()
