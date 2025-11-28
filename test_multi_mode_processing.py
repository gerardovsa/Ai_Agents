"""
Test Multi-Mode File Processing

Tests all 4 new processing modes:
1. extract - Markdown text extraction (baseline)
2. convert_pdf - Convert office docs to PDF
3. convert_image - Convert to PNG/JPEG images
4. hybrid - Text + images combined

Usage:
    python test_multi_mode_processing.py
"""

import sys
import os
sys.path.insert(0, os.path.dirname(__file__))

from AI_infrastructure.core.universal_file_handler import UniversalFileHandler
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
import io


def create_test_docx():
    """Create a simple test DOCX file"""
    doc = Document()
    doc.add_heading('Test Business Report', 0)
    doc.add_paragraph('This is a test document for multi-mode processing.')
    
    doc.add_heading('Section 1: Overview', 1)
    doc.add_paragraph('This section contains an overview of the test.')
    
    doc.add_heading('Section 2: Metrics', 1)
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = 'Metric'
    table.cell(0, 1).text = 'Value'
    table.cell(1, 0).text = 'Revenue'
    table.cell(1, 1).text = '$1.5M'
    table.cell(2, 0).text = 'Growth'
    table.cell(2, 1).text = '+15%'
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def create_test_xlsx():
    """Create a simple test XLSX file"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Data"
    
    ws['A1'] = 'Product'
    ws['B1'] = 'Q1'
    ws['C1'] = 'Q2'
    
    ws['A2'] = 'Product A'
    ws['B2'] = 1000
    ws['C2'] = 1200
    
    ws['A3'] = 'Product B'
    ws['B3'] = 800
    ws['C3'] = 950
    
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


def create_test_pptx():
    """Create a simple test PPTX file"""
    prs = Presentation()
    
    # Slide 1
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Q4 Results"
    
    # Slide 2
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Key Metrics"
    slide2.placeholders[1].text = "Revenue: $2.5M\nGrowth: +15%\nProfit: $500K"
    
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer.read()


def test_mode(handler, file_data, filename, content_type, mode, **kwargs):
    """Test a specific processing mode"""
    print(f"\n{'='*60}")
    print(f"Testing mode: {mode}")
    print(f"File: {filename}")
    print('='*60)
    
    try:
        result = handler.process_file(
            source='bytes',
            source_id={
                'filename': filename,
                'content_type': content_type,
                'data': file_data
            },
            mode=mode,
            **kwargs
        )
        
        if result['success']:
            print(f"[OK] Mode: {result['method']}")
            print(f"[OK] Token estimate: {result['metadata']['token_estimate']:,}")
            
            if result.get('content_block'):
                print(f"[OK] Single content block created")
                if result['content_block']['type'] == 'text':
                    preview = result['content_block']['text'][:200]
                    print(f"[OK] Text preview: {preview}...")
                elif result['content_block']['type'] == 'image':
                    print(f"[OK] Image content block (base64)")
                elif result['content_block']['type'] == 'document':
                    print(f"[OK] Document content block (PDF)")
            
            if result.get('content_blocks'):
                print(f"[OK] Multiple content blocks: {len(result['content_blocks'])}")
                for i, block in enumerate(result['content_blocks'], 1):
                    print(f"    - Block {i}: {block['type']}")
            
            if result['metadata'].get('extraction_metadata'):
                em = result['metadata']['extraction_metadata']
                print(f"[OK] Extraction metadata:")
                print(f"    - Format: {em.get('output_format')}")
                print(f"    - Words: {em.get('word_count'):,}")
            
            if result['metadata'].get('conversion_metadata'):
                cm = result['metadata']['conversion_metadata']
                print(f"[OK] Conversion metadata:")
                print(f"    - Method: {cm.get('conversion_method')}")
                if cm.get('image_count'):
                    print(f"    - Images: {cm.get('image_count')}")
            
            return True
        else:
            print(f"[FAIL] Error: {result.get('error')}")
            return False
    
    except Exception as e:
        print(f"[FAIL] Exception: {str(e)}")
        import traceback
        traceback.print_exc()
        return False


def main():
    print("\n" + "="*60)
    print("MULTI-MODE FILE PROCESSING TEST SUITE")
    print("="*60)
    
    # Initialize handler
    handler = UniversalFileHandler(user_id=1)
    
    # Create test files
    print("\n[INFO] Creating test files...")
    docx_data = create_test_docx()
    print("[OK] DOCX created (test_report.docx)")
    
    xlsx_data = create_test_xlsx()
    print("[OK] XLSX created (test_sales.xlsx)")
    
    pptx_data = create_test_pptx()
    print("[OK] PPTX created (test_presentation.pptx)")
    
    # Test results
    results = {
        'docx': {},
        'xlsx': {},
        'pptx': {}
    }
    
    # ========== TEST DOCX ==========
    
    print("\n" + "#"*60)
    print("# TESTING DOCX FILE")
    print("#"*60)
    
    # Test 1: Extract mode (baseline)
    results['docx']['extract'] = test_mode(
        handler, docx_data, 'test_report.docx',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'extract'
    )
    
    # Test 2: Convert to PDF
    results['docx']['convert_pdf'] = test_mode(
        handler, docx_data, 'test_report.docx',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'convert_pdf'
    )
    
    # Test 3: Convert to images (PNG)
    results['docx']['convert_image_png'] = test_mode(
        handler, docx_data, 'test_report.docx',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'convert_image',
        image_format='png',
        image_dpi=150
    )
    
    # Test 4: Convert to images (JPEG)
    results['docx']['convert_image_jpeg'] = test_mode(
        handler, docx_data, 'test_report.docx',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'convert_image',
        image_format='jpeg',
        image_dpi=120
    )
    
    # Test 5: Hybrid mode
    results['docx']['hybrid'] = test_mode(
        handler, docx_data, 'test_report.docx',
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        'hybrid'
    )
    
    # ========== TEST XLSX ==========
    
    print("\n" + "#"*60)
    print("# TESTING XLSX FILE")
    print("#"*60)
    
    results['xlsx']['extract'] = test_mode(
        handler, xlsx_data, 'test_sales.xlsx',
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'extract'
    )
    
    results['xlsx']['convert_pdf'] = test_mode(
        handler, xlsx_data, 'test_sales.xlsx',
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'convert_pdf'
    )
    
    results['xlsx']['convert_image'] = test_mode(
        handler, xlsx_data, 'test_sales.xlsx',
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'convert_image',
        image_format='png'
    )
    
    results['xlsx']['hybrid'] = test_mode(
        handler, xlsx_data, 'test_sales.xlsx',
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        'hybrid'
    )
    
    # ========== TEST PPTX ==========
    
    print("\n" + "#"*60)
    print("# TESTING PPTX FILE")
    print("#"*60)
    
    results['pptx']['extract'] = test_mode(
        handler, pptx_data, 'test_presentation.pptx',
        'application/vnd.openxmlformats-officedocument.presentationml.presentation',
        'extract'
    )
    
    results['pptx']['convert_pdf'] = test_mode(
        handler, pptx_data, 'test_presentation.pptx',
        'application/vnd.openxmlformats-officedocument.presentationml.presentation',
        'convert_pdf'
    )
    
    results['pptx']['convert_image'] = test_mode(
        handler, pptx_data, 'test_presentation.pptx',
        'application/vnd.openxmlformats-officedocument.presentationml.presentation',
        'convert_image',
        image_format='jpeg'
    )
    
    results['pptx']['hybrid'] = test_mode(
        handler, pptx_data, 'test_presentation.pptx',
        'application/vnd.openxmlformats-officedocument.presentationml.presentation',
        'hybrid'
    )
    
    # ========== SUMMARY ==========
    
    print("\n" + "="*60)
    print("TEST SUMMARY")
    print("="*60)
    
    total_tests = 0
    passed_tests = 0
    
    for file_type, modes in results.items():
        print(f"\n{file_type.upper()} Results:")
        for mode, success in modes.items():
            status = "[OK]" if success else "[FAIL]"
            print(f"  {status} {mode}")
            total_tests += 1
            if success:
                passed_tests += 1
    
    print(f"\n{'='*60}")
    print(f"TOTAL: {passed_tests}/{total_tests} tests passed")
    print(f"{'='*60}")
    
    if passed_tests == total_tests:
        print("\n[SUCCESS] All multi-mode processing tests PASSED!")
        return 0
    else:
        print(f"\n[PARTIAL] {passed_tests}/{total_tests} tests passed")
        return 1


if __name__ == '__main__':
    sys.exit(main())
